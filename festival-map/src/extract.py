import sys, json
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
out=[]
for i,s in enumerate(p.slides):
    shapes=[]
    def walk(shs):
        for sh in shs:
            if sh.shape_type==6: walk(sh.shapes); continue
            fill=None
            try:
                if sh.fill.type==1:
                    try: fill=str(sh.fill.fore_color.rgb)
                    except Exception: fill='theme:'+str(sh.fill.fore_color.theme_color)
            except Exception: pass
            geom=None
            try: geom=sh.auto_shape_type and str(sh.auto_shape_type)
            except Exception: pass
            shapes.append(dict(name=sh.name,x=sh.left,y=sh.top,w=sh.width,h=sh.height,rot=sh.rotation,fill=fill,geom=geom,text=sh.text_frame.text if sh.has_text_frame else ''))
    walk(s.shapes)
    out.append(shapes)
json.dump(out,open(sys.argv[2],'w'),ensure_ascii=False)
for i,sl in enumerate(out):
    from collections import Counter
    print(i+1, len(sl), Counter((x['fill'],x['geom']) for x in sl))
